"""Generate legitimate documents that use white-on-dark designs, to check for false positives."""
import os
import sys

import docx
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import openpyxl
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor as PRGB
import pymupdf


def clean_docx(path):
    d = docx.Document()
    d.add_heading("Team Budget", 1)
    d.add_paragraph("The following table summarizes spending for the quarter.")
    t = d.add_table(rows=2, cols=2)
    for i, txt in enumerate(("Item", "Cost")):
        cell = t.rows[0].cells[i]
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "1F3864")
        cell._tc.get_or_add_tcPr().append(shd)
        r = cell.paragraphs[0].add_run(txt)
        r.font.color.rgb = RGBColor(255, 255, 255)
    t.rows[1].cells[0].text = "Laptops"
    t.rows[1].cells[1].text = "$4,000"
    p = d.add_paragraph()
    r = p.add_run("Highlighted white text on a dark highlight")
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.highlight_color = 2  # dark blue? WD_COLOR_INDEX.BLUE
    d.save(path)


def clean_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, h in enumerate(("Region", "Sales"), 1):
        c = ws.cell(1, i, h)
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = PatternFill("solid", fgColor="305496")
    ws.append(["North", 100])
    ws.append(["South", 200])
    ws["D1"] = "Total"
    ws["E1"] = "=SUM(B2:B3)"
    wb.save(path)


def clean_pptx(path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = PRGB(0x10, 0x10, 0x30)
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "White title on a dark slide background"
    r.font.color.rgb = PRGB(255, 255, 255)
    box = s.shapes.add_shape(1, Inches(1), Inches(3), Inches(4), Inches(1))
    box.fill.solid()
    box.fill.fore_color.rgb = PRGB(0xC0, 0, 0)
    box.text_frame.text = "Callout in a red box"
    box.text_frame.paragraphs[0].runs[0].font.color.rgb = PRGB(255, 255, 255)
    prs.save(path)


def clean_pdf(path):
    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect(pymupdf.Rect(0, 0, page.rect.width, 80), color=None, fill=(0.1, 0.2, 0.4))
    page.insert_text((72, 50), "Company Newsletter - white header on navy banner", fontsize=18, color=(1, 1, 1))
    page.insert_text((72, 120), "Welcome to this month's edition. Sales grew 12 percent.", fontsize=11)
    doc.save(path)


if __name__ == "__main__":
    out = sys.argv[1]
    os.makedirs(out, exist_ok=True)
    clean_docx(os.path.join(out, "clean.docx"))
    clean_xlsx(os.path.join(out, "clean.xlsx"))
    clean_pptx(os.path.join(out, "clean.pptx"))
    clean_pdf(os.path.join(out, "clean.pdf"))
    print("ok")
